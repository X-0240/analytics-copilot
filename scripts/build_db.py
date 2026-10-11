"""把 UCI Online Retail 建成只读分析库（DuckDB）并导出元数据。

顺序（全部可复现，不需要联网）：
  1. 读 `data/raw/Online Retail.xlsx` → 表 `raw_online_retail`（原样，8 列）；
  2. 建视图 `v_sales`（加口径列：is_return / line_revenue / invoice_ts / is_cancelled）；
  3. 建视图 `v_returns`（只看退货行）；
  4. 导出 `data/metadata.json`：对象清单、列信息、空值率、口径说明、白名单。

设计约束：应用层只用**只读连接**打开这个库，并且只暴露白名单里的视图；
原始表不出现在白名单里（防止绕过口径列直接查）。

用法：
    python scripts/build_db.py            # 建库（幂等：先删再建）
    python scripts/build_db.py --check    # 只核对已有库与元数据
"""

import argparse
import json
import sys
from pathlib import Path

import duckdb

ROOT = Path(__file__).resolve().parents[1]
RAW_XLSX = ROOT / "data" / "raw" / "Online Retail.xlsx"
RAW_CSV = ROOT / "data" / "raw" / "online_retail.csv"
DB_PATH = ROOT / "data" / "analytics.duckdb"
META_PATH = ROOT / "data" / "metadata.json"

# 应用层可见的对象白名单：只给视图，不给原始表
ALLOWED_VIEWS = ("v_sales", "v_returns")

VIEW_SALES = """
CREATE OR REPLACE VIEW v_sales AS
SELECT
    invoice_no,
    stock_code,
    description,
    CAST(quantity AS BIGINT)      AS quantity,
    CAST(invoice_ts AS TIMESTAMP) AS invoice_ts,
    CAST(unit_price AS DOUBLE)   AS unit_price,
    CAST(customer_id AS BIGINT)  AS customer_id,
    country,
    (quantity < 0)                  AS is_return,
    (LEFT(invoice_no, 1) = 'C')     AS is_cancelled,
    (quantity * unit_price)         AS line_revenue
FROM raw_online_retail
"""

VIEW_RETURNS = """
CREATE OR REPLACE VIEW v_returns AS
SELECT * FROM v_sales WHERE is_return
"""

CALIBER = {
    "grain": "一行 = 一张发票里的一行商品（invoice line），不是订单级",
    "line_revenue": "quantity * unit_price，**含退货为负**；做汇总时不要再次过滤符号",
    "is_return": "quantity < 0；退货不计入正向销售额时可加 WHERE NOT is_return",
    "is_cancelled": "发票号以 C 开头（取消单），与 is_return 高度相关但不完全等同",
    "customer_id": "缺失率约 25%，**客户维度聚合前必须显式说明样本量**",
    "time_range": "2010-12-01 ~ 2011-12-09（约 1 年），跨年比较要考虑季节性",
    "no_cost_field": "数据集没有成本/毛利字段，因此**不能做利润归因**",
    "country": "38 个国家/地区，'United Kingdom' 占绝大多数，跨国比较要报占比",
}


def load_raw(con, try_ext=False):
    """读 Excel 进原始表。

    先试 DuckDB 的 excel 扩展；本机实测该扩展需要联网下载，在代理环境下返回
    HTTP 302 而失败（见 metadata.note），因此默认路径是：
    openpyxl 流式转 CSV（一次）→ DuckDB 读 CSV（很快）。

    `try_ext=True` 时才尝试扩展（本机每次尝试都要等代理超时约 5 分钟，故默认关闭）。
    """
    note = ""
    if try_ext:
        try:
            con.execute("INSTALL excel; LOAD excel;")
            con.execute(
                "CREATE OR REPLACE TABLE raw_online_retail AS "
                "SELECT * FROM read_xlsx(%s, header=true)" % _lit(str(RAW_XLSX))
            )
            return "duckdb_excel_extension", note
        except Exception as exc:
            note = "excel 扩展不可用（%s），改用 openpyxl→CSV→read_csv_auto" % (
                type(exc).__name__
            )
            print("[warn] %s" % note)
    else:
        note = (
            "默认不走 DuckDB excel 扩展：本机代理下该扩展下载失败"
            "（HTTP 302），改用 openpyxl→CSV→read_csv_auto"
        )
    return _load_via_csv(con), note


def _lit(text):
    """把路径写成 SQL 字符串字面量（只用于本地固定路径）。"""
    return "'" + text.replace("'", "''") + "'"


def _xlsx_to_csv():
    """openpyxl 只读模式流式写 CSV（只在 CSV 不存在时执行一次）。"""
    import csv

    from openpyxl import load_workbook

    if RAW_CSV.exists():
        print("[info] 复用已有 CSV：%s" % RAW_CSV.name)
        return
    book = load_workbook(RAW_XLSX, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(name).strip().lower().replace(" ", "_") for name in next(rows)]
    total = 0
    with RAW_CSV.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for row in rows:
            writer.writerow(["" if value is None else value for value in row])
            total += 1
    book.close()
    print("[info] 已转换 CSV：%d 行" % total)


def _load_via_csv(con):
    """CSV → DuckDB 原始表（列名与类型显式声明，避免自动推断踩坑）。"""
    _xlsx_to_csv()
    # CSV 表头来自 Excel 原列名（InvoiceNo→invoiceno 等），在这里统一改成库里
    # 视图使用的下划线命名；同时显式声明类型，不依赖自动推断。
    con.execute(
        """
        CREATE OR REPLACE TABLE raw_online_retail AS
        SELECT
            invoiceno   AS invoice_no,
            stockcode   AS stock_code,
            description AS description,
            quantity    AS quantity,
            invoicedate AS invoice_ts,
            unitprice   AS unit_price,
            customerid  AS customer_id,
            country     AS country
        FROM read_csv_auto(
            %s, header=true, null_padding=true,
            types={
                'invoiceno': 'VARCHAR', 'stockcode': 'VARCHAR',
                'description': 'VARCHAR', 'quantity': 'BIGINT',
                'invoicedate': 'TIMESTAMP', 'unitprice': 'DOUBLE',
                'customerid': 'BIGINT', 'country': 'VARCHAR'
            }
        )
        """
        % _lit(str(RAW_CSV))
    )
    return "openpyxl_csv"


def build_metadata(con, source, note=""):
    """导出对象清单、列信息、空值率与口径说明。"""
    # DuckDB 的 SHOW TABLES 只返回一列，这里用 information_schema 拿名称与类型
    tables = [
        {"name": row[0], "type": row[1]}
        for row in con.execute(
            "SELECT table_name, table_type FROM information_schema.tables "
            "WHERE table_schema = 'main' ORDER BY table_name"
        ).fetchall()
    ]
    columns = {}
    for entry in tables:  # tables 是 dict 列表，直接解包会取到键名
        name = entry["name"]
        columns[name] = [
            {"column": row[0], "dtype": row[1]}
            for row in con.execute('DESCRIBE "%s"' % name).fetchall()
        ]
    stats = con.execute(
        """
        SELECT COUNT(*)                                   AS rows,
               MIN(invoice_ts)                            AS ts_min,
               MAX(invoice_ts)                            AS ts_max,
               COUNT(DISTINCT country)                    AS countries,
               COUNT(DISTINCT stock_code)                 AS stock_codes,
               COUNT(DISTINCT customer_id)                AS customers,
               SUM(CASE WHEN customer_id IS NULL THEN 1 ELSE 0 END) AS missing_customer,
               SUM(CASE WHEN is_return THEN 1 ELSE 0 END) AS return_lines,
               ROUND(SUM(line_revenue), 2)                AS revenue_total
        FROM v_sales
        """
    ).fetchone()
    return {
        "source_file": RAW_XLSX.name,
        "built_by": source,
        "note": note,
        "objects": tables,
        "columns": columns,
        "allowed_views": list(ALLOWED_VIEWS),
        "profile": {
            "rows": stats[0],
            "ts_min": str(stats[1]),
            "ts_max": str(stats[2]),
            "countries": stats[3],
            "stock_codes": stats[4],
            "customers": stats[5],
            "missing_customer_rows": stats[6],
            "missing_customer_ratio": round(stats[6] / stats[0], 4),
            "return_lines": stats[7],
            "revenue_total": stats[8],
        },
        "caliber": CALIBER,
    }


def main():
    parser = argparse.ArgumentParser(prog="build_db")
    parser.add_argument("--check", action="store_true", help="只核对已有库")
    parser.add_argument(
        "--try-excel-ext",
        action="store_true",
        help="尝试 DuckDB excel 扩展（本机代理下会等约 5 分钟才失败）",
    )
    args = parser.parse_args()

    if not RAW_XLSX.exists():
        raise SystemExit("缺少原始数据：%s" % RAW_XLSX)
    if args.check:
        if not DB_PATH.exists() or not META_PATH.exists():
            raise SystemExit("库或元数据不存在，先建库")
        con = duckdb.connect(str(DB_PATH), read_only=True)
        meta = json.loads(META_PATH.read_text(encoding="utf-8"))
        print("库 %s ｜ 对象 %s" % (DB_PATH.name, meta["allowed_views"]))
        print("行数 %s ｜ 口径项 %d 条" % (
            meta["profile"]["rows"], len(meta["caliber"])
        ))
        missing = [
            name for name in meta["allowed_views"]
            if name
            not in {
                row[0]
                for row in con.execute(
                    "SELECT table_name FROM information_schema.tables "
                    "WHERE table_schema = 'main'"
                ).fetchall()
            }
        ]
        if missing:
            raise SystemExit("白名单对象缺失：%s" % missing)
        print("核对通过")
        con.close()
        return

    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    if DB_PATH.exists():
        DB_PATH.unlink()
    con = duckdb.connect(str(DB_PATH))
    source, note = load_raw(con, try_ext=args.try_excel_ext)
    con.execute(VIEW_SALES)
    con.execute(VIEW_RETURNS)
    meta = build_metadata(con, source, note)
    META_PATH.write_text(
        json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    con.close()
    print("建库完成 → %s（%s）" % (DB_PATH, source))
    print("白名单视图 %s ｜ 行数 %s ｜ 缺 customer_id %.2f%%" % (
        meta["allowed_views"],
        meta["profile"]["rows"],
        meta["profile"]["missing_customer_ratio"] * 100,
    ))


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
